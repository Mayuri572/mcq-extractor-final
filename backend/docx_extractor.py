"""
docx_extractor.py
Extracts plain text (in document order, including simple tables) from a
.docx file using python-docx.
"""

from docx import Document
from docx.oxml.ns import qn


def extract_docx_text(path: str) -> str:
    doc = Document(path)
    lines = []

    # Iterate the document body in order so paragraphs and tables interleave
    # the way they visually appear, rather than all paragraphs then all tables.
    body = doc.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            para = _find_paragraph(doc, child)
            if para is not None:
                text = para.text.strip()
                if text:
                    lines.append(text)
        elif child.tag == qn("w:tbl"):
            table = _find_table(doc, child)
            if table is not None:
                for row in table.rows:
                    for cell in row.cells:
                        text = cell.text.strip()
                        if text:
                            lines.append(text)

    return "\n".join(lines)


def _find_paragraph(doc, element):
    for para in doc.paragraphs:
        if para._p is element:
            return para
    return None


def _find_table(doc, element):
    for table in doc.tables:
        if table._tbl is element:
            return table
    return None
