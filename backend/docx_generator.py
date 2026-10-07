"""
docx_generator.py
Writes the extracted MCQs into a simple, professionally formatted Word
document using python-docx.
"""

import os

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

OPTION_LETTERS = ["A", "B", "C", "D"]


def write_docx_file(result: dict, output_dir: str, job_id: str) -> str:
    os.makedirs(output_dir, exist_ok=True)
    out_path = os.path.join(output_dir, f"{job_id}.docx")

    doc = Document()

    for section in doc.sections:
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)

    title = doc.add_heading("Extracted MCQs", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT

    meta = doc.add_paragraph()
    meta_run = meta.add_run(
        f"Source: {result.get('document_name', '-')}   |   "
        f"Format: {result.get('format', '-')}   |   "
        f"Language(s): {', '.join(result.get('languages', [])) or '-'}   |   "
        f"Total Questions: {result.get('total_questions', 0)}"
    )
    meta_run.italic = True
    meta_run.font.size = Pt(10)

    doc.add_paragraph()

    for q in result.get("questions", []):
        q_para = doc.add_paragraph()
        q_run = q_para.add_run(f"Q{q['question_number']}. {q['question']}")
        q_run.bold = True
        q_run.font.size = Pt(12)

        options = q.get("options", {})
        for letter in OPTION_LETTERS:
            if letter in options:
                opt_para = doc.add_paragraph()
                opt_para.paragraph_format.left_indent = Cm(0.8)
                opt_para.add_run(f"{letter}. {options[letter]}").font.size = Pt(11)

        doc.add_paragraph()  # spacing between questions

    incomplete_questions = result.get("incomplete_questions", [])
    _write_flagged_section(
        doc, incomplete_questions,
        heading="Incomplete Questions",
        note=("Found but could not be fully extracted (missing question text or "
              "one or more options); kept separate from the list above."),
        detail_fn=lambda q: f"Missing: {', '.join(q.get('missing', [])) or 'unknown'}",
    )

    needs_review_questions = result.get("needs_review_questions", [])
    _write_flagged_section(
        doc, needs_review_questions,
        heading="Needs Review",
        note=("The parser could not confidently determine the question/option "
              "boundaries for these (e.g. no question marker was detected, or "
              "an option letter appeared twice, which usually means a question "
              "marker was missed in the source) - please verify manually rather "
              "than trusting this content as-is."),
        detail_fn=lambda q: f"Reason: {', '.join(q.get('reasons', [])) or 'unknown'}",
    )

    doc.save(out_path)
    return out_path


def _write_flagged_section(doc, items, heading, note, detail_fn):
    """Shared writer for the Incomplete Questions / Needs Review sections
    (same layout, different heading/note/per-item detail line)."""
    if not items:
        return

    doc.add_paragraph()
    doc.add_heading(heading, level=2)
    note_para = doc.add_paragraph()
    note_run = note_para.add_run(note)
    note_run.italic = True
    note_run.font.size = Pt(10)

    for i, q in enumerate(items, start=1):
        q_para = doc.add_paragraph()
        q_run = q_para.add_run(f"[{i}] {q.get('question') or '(no question text found)'}")
        q_run.bold = True
        q_run.font.size = Pt(12)

        detail_para = doc.add_paragraph()
        detail_run = detail_para.add_run(detail_fn(q))
        detail_run.italic = True
        detail_run.font.size = Pt(10)

        options = q.get("options", {})
        for letter in OPTION_LETTERS:
            value = options.get(letter, "")
            if value:
                opt_para = doc.add_paragraph()
                opt_para.paragraph_format.left_indent = Cm(0.8)
                opt_para.add_run(f"{letter}. {value}").font.size = Pt(11)

        doc.add_paragraph()
